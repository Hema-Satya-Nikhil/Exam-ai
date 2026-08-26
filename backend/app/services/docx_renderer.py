from typing import Any, Dict
from docx import Document
from docx.shared import PT, Inches
import io
from ..models.paper import Paper

class DOCXRenderer:
    """
    Handles DOCX generation from structured paper content
    """
    def __init__(self, content: Dict[str, Any]):
        self.content = content

    def generate(self) -> bytes:
        """
        Generate DOCX bytes from the paper content
        """
        buffer = io.BytesIO()
        doc = Document()
        
        # Add institution header
        if 'institution' in self.content:
            doc.add_heading(self.content['institution'], level=1)

        # Add department
        if 'department' in self.content:
            doc.add_heading(self.content['department'], level=2)

        # Add examination details
        if 'examination' in self.content:
            examination = doc.add_paragraph()
            run = examination.add_run(f"{self.content['examination']['name']} - {self.content['examination']['year']}")
            run.bold = True
            run.font.size = Pt(14)

        # Add subject details
        if 'subject' in self.content:
            subject_section = doc.add_heading(self.content['subject']['name'], level=3)
            subject_section.add_run(f"Subject Code: {self.content['subject']['code']}").font.size = Pt(12)
            subject_section.add_run(f"Duration: {self.content['subject']['duration']} minutes").font.size = Pt(12)

        # Add instructions
        if 'instructions' in self.content:
            instructions = doc.add_heading("INSTRUCTIONS", level=4)
            for instruction in self.content['instructions']:
                instructions.add_paragraph(instruction, style="List Bullet")

        # Add sections and questions
        if 'sections' in self.content:
            for section in self.content['sections']:
                section_header = doc.add_heading(f"{section['title']} (Total Marks: {section['total_marks']})", level=4)
                for question in section['questions']:
                    question_para = section_header.add_paragraph()
                    question_para.add_run(f"{question['number}]. {question['text']} [Marks: {question['marks']}"])
                    question_para.runs[0].font.bold = True
                    question_para.runs[0].font.size = Pt(11)
                    
                    if 'choices' in question:
                        for choice in question['choices']:
                            choice_para = question_para.add_run(f"- {choice}")
                            choice_para.font.size = Pt(10)

        doc.save(buffer)
        return buffer.getvalue()

    def render(self, content: str) -> bytes:
        """
        Render dynamic content to DOCX
        """
        doc = Document()
        doc.add_paragraph(content)
        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()