from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


class StructuredPaperRenderer:
    def render_pdf(self, paper_json: dict, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(str(output_path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(name="TitleGlass", parent=styles["Title"], fontSize=16, leading=20, alignment=1, spaceAfter=10)
        heading_style = ParagraphStyle(name="HeadingGlass", parent=styles["Heading2"], fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
        body_style = ParagraphStyle(name="BodyGlass", parent=styles["BodyText"], fontSize=9.5, leading=13)

        story = []
        story.append(Paragraph(paper_json.get("institution_name", "Question Paper"), title_style))
        story.append(Paragraph(paper_json.get("department_name", "Department"), body_style))
        story.append(Spacer(1, 6))
        story.append(Paragraph(f"<b>{paper_json.get('exam_name', 'Examination')}</b>", heading_style))
        story.append(Paragraph(f"Subject: {paper_json.get('subject_name', 'Subject')} {paper_json.get('subject_code', '')}", body_style))
        story.append(Paragraph(f"Duration: {paper_json.get('duration_minutes', 0)} minutes | Maximum Marks: {paper_json.get('total_marks', 0)}", body_style))
        story.append(Spacer(1, 8))

        for instruction in paper_json.get("instructions", []):
            story.append(Paragraph(f"• {instruction}", body_style))
        if paper_json.get("instructions"):
            story.append(Spacer(1, 6))

        for section in paper_json.get("sections", []):
            story.append(Paragraph(section.get("name", "Section"), heading_style))
            if section.get("instructions"):
                story.append(Paragraph(section["instructions"], body_style))
            questions = section.get("questions", [])
            table_data = [["Q No.", "Marks", "Question"]]
            for question in questions:
                question_text = question.get("question_text", "")
                choice_group = question.get("choice_group")
                if choice_group:
                    question_text = f"{question_text} [{choice_group}]"
                table_data.append([
                    str(question.get("question_number", "")),
                    str(question.get("marks", "")),
                    Paragraph(question_text, body_style),
                ])
            table = Table(table_data, colWidths=[16 * mm, 18 * mm, 145 * mm], repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 8))

        doc.build(story)
        return output_path

    def render_docx(self, paper_json: dict, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.core_properties.title = paper_json.get("title", "Question Paper")
        document.core_properties.subject = paper_json.get("subject_name", "Subject")

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(paper_json.get("institution_name", "Question Paper"))
        run.bold = True
        run.font.size = Pt(16)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.add_run(paper_json.get("department_name", "Department"))
        subtitle_run.font.size = Pt(11)

        document.add_paragraph(f"Examination: {paper_json.get('exam_name', 'Examination')}")
        document.add_paragraph(f"Subject: {paper_json.get('subject_name', 'Subject')} {paper_json.get('subject_code', '')}")
        document.add_paragraph(f"Duration: {paper_json.get('duration_minutes', 0)} minutes")
        document.add_paragraph(f"Maximum Marks: {paper_json.get('total_marks', 0)}")

        if paper_json.get("instructions"):
            document.add_paragraph("Instructions", style="Heading 1")
            for instruction in paper_json["instructions"]:
                document.add_paragraph(instruction, style="List Bullet")

        for section in paper_json.get("sections", []):
            document.add_paragraph(section.get("name", "Section"), style="Heading 2")
            if section.get("instructions"):
                document.add_paragraph(section["instructions"])
            for question in section.get("questions", []):
                line = f"Q{question.get('question_number', '')}. ({question.get('marks', '')} marks) {question.get('question_text', '')}"
                if question.get("choice_group"):
                    line += f" [{question['choice_group']}]"
                document.add_paragraph(line)

        document.save(str(output_path))
        return output_path
