"""Issue 3 regression: official PDF/DOCX documents must not leak internal metadata.

Renders REAL documents through ``StructuredPaperRenderer`` and extracts the
actual text (pypdf for PDF, python-docx for DOCX), asserting that no internal
implementation metadata reaches the official exam output while valid exam
content (subject, sections, numbering, marks, question wording) is preserved.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.services.documents.renderer import StructuredPaperRenderer

FORBIDDEN_MARKERS = (
    "[choice_group",
    "[cg-",
    "choice_group_id",
    "choice_member_id",
    "source_references",
    "section_blah",
    "Bloom:",
    "Difficulty:",
    "Topic: ",
    "Section C - Unit",
    '{"',
)

VALID_MARKERS = ("Section A", "Section B", "Analyze the appropriateness")


def _legacy_paper_json() -> dict:
    """Legacy flat paper with metadata present on the question dicts.

    The renderer must print ONLY question_number, marks and question_text —
    never the surrounding metadata fields.
    """
    return {
        "institution_name": "TEST INSTITUTE",
        "department_name": "CSE",
        "exam_name": "MID 1",
        "subject_name": "Operating Systems",
        "subject_code": "CS501",
        "duration_minutes": 110,
        "total_marks": 50,
        "instructions": ["Answer all questions."],
        "sections": [
            {
                "name": "Section A",
                "instructions": "Answer all.",
                "questions": [
                    {
                        "question_number": 1,
                        "marks": 2,
                        "question_text": "Define a process.",
                        "unit": 1,
                        "topic": "Process Concept",
                        "bloom": "L2",
                        "difficulty": "easy",
                        "question_type": "descriptive",
                        "source_references": ["src-1"],
                    },
                    {
                        "question_number": 2,
                        "marks": 5,
                        "question_text": "Analyze the appropriateness of various CPU scheduling algorithms.",
                        "unit": 2,
                        "topic": "Scheduling criteria",
                        "bloom": "L5",
                        "difficulty": "hard",
                        "question_type": "analytical",
                        "choice_group_id": "cg-56",
                        "choice_member_id": "cm-1",
                    },
                ],
            },
            {
                "name": "Section B",
                "instructions": "",
                "questions": [
                    {
                        "question_number": 3,
                        "marks": 5,
                        "question_text": "Explain multithreading models.",
                        "unit": 2,
                        "topic": "Multithreading models",
                        "bloom": "L3",
                        "difficulty": "medium",
                    }
                ],
            },
        ],
    }


def _composite_compat_paper_json() -> dict:
    """Composite exam (compat flat sections view) incl. OR alternatives."""
    return {
        "institution_name": "TEST INSTITUTE",
        "department_name": "CSE",
        "exam_name": "MID 1",
        "subject_name": "Data Warehousing and Data Mining",
        "subject_code": "CS601",
        "duration_minutes": 110,
        "total_marks": 50,
        "instructions": [],
        "sections": [
            {
                "name": "Part A",
                "instructions": "10 x 2 Marks",
                "questions": [
                    {
                        "question_number": 1,
                        "marks": 2,
                        "question_text": "List the components of an operating system.",
                        "exam_part": "short_answer",
                        "unit": 1,
                        "topic": "Process Concept",
                        "bloom": "L2",
                    }
                ],
            },
            {
                "name": "Part B",
                "instructions": "30 Marks",
                "questions": [
                    {
                        "question_number": 2,
                        "marks": 5,
                        "question_text": "Explain process scheduling.",
                        "exam_part": "main",
                        "group_number": 2,
                        "part_label": "a",
                        "choice_group_id": "cg1",
                        "choice_member_id": "cm-a",
                    },
                    {
                        "question_number": 3,
                        "marks": 5,
                        "question_text": "Explain inter-process communication.",
                        "exam_part": "main",
                        "group_number": 2,
                        "part_label": "b",
                        "choice_group_id": "cg1",
                        "choice_member_id": "cm-b",
                    },
                ],
            },
        ],
    }


def _pdf_text(path: Path) -> str:
    reader = PdfReader(io.BytesIO(path.read_bytes()))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.parametrize("key", ["legacy_pdf", "legacy_docx"])
def test_legacy_content_preserved(rendered, key):
    text = rendered[key]
    for marker in VALID_MARKERS:
        assert marker in text, f"missing valid content {marker!r} in {key}"
    assert "Define a process." in text
    assert "Operating Systems" in text


@pytest.mark.parametrize("key", ["comp_pdf", "comp_docx"])
def test_composite_content_preserved(rendered, key):
    text = rendered[key]
    assert "Part A" in text
    assert "Part B" in text
    assert "Explain process scheduling." in text
    assert "cg1" not in text and "cm-a" not in text


def test_renderer_never_adds_metadata_prefixes(rendered):
    """Renderer must not prepend metadata like '(5 marks) Section C - Unit 2'."""
    for key in ("legacy_pdf", "legacy_docx", "comp_pdf", "comp_docx"):
        for line in rendered[key].splitlines():
            stripped = line.strip()
            assert "Section C -" not in stripped, stripped



@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    out = tmp_path_factory.mktemp("exports")
    renderer = StructuredPaperRenderer()
    legacy_pdf = renderer.render_pdf(_legacy_paper_json(), out / "legacy.pdf")
    legacy_docx = renderer.render_docx(_legacy_paper_json(), out / "legacy.docx")
    comp_pdf = renderer.render_pdf(_composite_compat_paper_json(), out / "comp.pdf")
    comp_docx = renderer.render_docx(_composite_compat_paper_json(), out / "comp.docx")
    return {
        "legacy_pdf": _pdf_text(legacy_pdf),
        "legacy_docx": _docx_text(legacy_docx),
        "comp_pdf": _pdf_text(comp_pdf),
        "comp_docx": _docx_text(comp_docx),
    }


@pytest.mark.parametrize("key", ["legacy_pdf", "legacy_docx", "comp_pdf", "comp_docx"])
def test_no_internal_metadata_in_document(rendered, key):
    text = rendered[key]
    for marker in FORBIDDEN_MARKERS:
        assert marker not in text, f"leaked {marker!r} in {key}:\n{text}"

